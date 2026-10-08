# -*- coding: utf-8 -*-
"""
Sinh file Excel "Bộ từ khoá Google Ads" cho 1 ZONE của Nha khoa Paris.

Phương pháp luận: MARKETING LẤY KHÁCH HÀNG LÀM TRUNG TÂM
    ZONE -> Chân dung KH -> Hành trình S1..S6 -> Cụm truy vấn ưu tiên (P1/P2/P3)
    -> Cụm nội dung + thống nhất mục tiêu -> Thực thi (nội dung / SEO-GEO / báo cáo)

Dùng:
    python tools/build_keyword_workbook.py zones/implant.json
    python tools/build_keyword_workbook.py "zones/*.json" -o out/

Cấu trúc file xuất ra (giữ đúng 5 sheet của template gốc):
    1. Tổng quan          - ZONE, chân dung KH, phân bổ ngân sách, mô hình phễu, nguyên tắc
    2. Bộ từ khoá         - 1 dòng = 1 từ khoá x kiểu khớp, gắn chân dung + chặng + rào cản
    3. Từ khoá phủ định   - danh sách chung + phủ định chéo điều hướng truy vấn
    4. Hành trình KH      - S1..S6, landing page, ngân hàng Hook x Format (FB/TikTok)
    5. Giá thầu & Tối ưu  - chấm điểm ưu tiên, ngưỡng, lộ trình 3 pha, A/B test, rủi ro
"""
import argparse
import glob
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- bảng màu
NAVY = "FF1E3A5E"
BAND = "FFDBE5F2"
ZEBRA = "FFF4F7FC"
WHITE = "FFFFFFFF"
TOTAL = "FFF2F2F2"
INPUT_BG = "FFFFF2CC"
INPUT_FG = "FF0000FF"
FORMULA_FG = "FF007F00"

FONT = "Arial"
VND = "#,##0"
PCT = "0%"
PCT1 = "0.0%"
NUM0 = "#,##0"
NUM2 = "0.00"
THIN = Side(style="thin", color="FFD9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _title(ws, row, text, size=14):
    c = ws.cell(row, 1, text)
    c.font = Font(name=FONT, size=size, bold=True, color=NAVY)
    return row + 1


def _note(ws, row, text):
    c = ws.cell(row, 1, text)
    c.font = Font(name=FONT, size=10, italic=True, color="FF667085")
    return row + 1


def _header(ws, row, labels):
    for i, lb in enumerate(labels, start=1):
        c = ws.cell(row, i, lb)
        c.font = Font(name=FONT, size=10, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
    ws.row_dimensions[row].height = 30
    return row + 1


def _band(ws, row, text, span=6):
    c = ws.cell(row, 1, text)
    c.font = Font(name=FONT, size=10, bold=True, color=NAVY)
    for i in range(1, span + 1):
        ws.cell(row, i).fill = PatternFill("solid", fgColor=BAND)
    return row + 1


def _cell(ws, row, col, value, fmt=None, kind=None, bold=False,
          fill=None, align=None, wrap=True):
    """kind: None | 'input' (nền vàng, chữ xanh) | 'formula' (chữ xanh lá)"""
    c = ws.cell(row, col, value)
    color = None
    if kind == "input":
        c.fill = PatternFill("solid", fgColor=INPUT_BG)
        color = INPUT_FG
    elif kind == "formula":
        color = FORMULA_FG
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    c.font = Font(name=FONT, size=10, bold=bold, color=color)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    c.border = BORDER
    return c


def _widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


# ============================================================ 1. Tổng quan
def sheet_tong_quan(wb, z):
    ws = wb.create_sheet("1. Tổng quan")
    _widths(ws, {"A": 34, "B": 20, "C": 46, "D": 34, "E": 22, "F": 22, "G": 12})
    zone = z["zone"]
    r = 1

    r = _title(ws, r, "BỘ TỪ KHOÁ GOOGLE ADS: " + zone["ten"].upper() + " | NHA KHOA PARIS")
    r = _note(ws, r, "Mục tiêu: " + zone["muc_tieu"])
    r = _note(ws, r, "Phương pháp luận lấy khách hàng làm trung tâm: ZONE -> chân dung KH "
                     "-> hành trình S1-S6 -> cụm truy vấn ưu tiên P1/P2/P3 -> cụm nội dung -> thực thi.")
    r = _note(ws, r, "Ô nền vàng chữ xanh là giả định đầu vào, thay bằng số thực của NKP. "
                     "Ô chữ xanh lá là công thức, không sửa tay.")
    r += 1

    # ---- A. chân dung khách hàng
    r = _title(ws, r, "A. CHÂN DUNG KHÁCH HÀNG - GỐC CỦA MỌI QUYẾT ĐỊNH PHÍA SAU", 12)
    r = _note(ws, r, "Không chọn từ khoá trước. Chọn người trước, rồi mới hỏi người đó gõ gì lên Google.")
    r = _header(ws, r, ["Mã", "Chân dung KH", "Insight / nỗi đau thật", "Rào cản chốt chính",
                        "Vào phễu ở chặng", "Chiến dịch phục vụ", "Giá trị ca"])
    for p in z["chan_dung"]:
        _cell(ws, r, 1, p["ma"], bold=True, align="center")
        _cell(ws, r, 2, p["ten"], bold=True)
        _cell(ws, r, 3, p["insight"])
        _cell(ws, r, 4, p["rao_can"])
        _cell(ws, r, 5, p["chang_vao"], align="center")
        _cell(ws, r, 6, p["chien_dich"])
        _cell(ws, r, 7, p["gia_tri_ca"], align="center")
        r += 1
    r += 1

    # ---- B. phân bổ ngân sách
    r = _title(ws, r, "B. PHÂN BỔ NGÂN SÁCH THEO CHIẾN DỊCH", 12)
    _cell(ws, r, 1, "Ngân sách Search/tháng (VND)", bold=True)
    budget_ref = "$B$" + str(r)
    _cell(ws, r, 2, zone["ngan_sach_thang"], fmt=VND, kind="input", align="center")
    _cell(ws, r, 3, "Giả định. Thay bằng ngân sách được duyệt.")
    r += 2

    r = _header(ws, r, ["Chiến dịch", "Ưu tiên", "Vai trò trong hành trình KH", "Tỷ trọng ngân sách",
                        "Ngân sách/tháng (VND)", "Ngân sách/ngày (VND)", "Số từ khoá"])
    cd_first = r
    cd_rows = []
    for cd in z["chien_dich"]:
        _cell(ws, r, 1, cd["ten"], bold=True)
        _cell(ws, r, 2, cd["uu_tien"], align="center")
        _cell(ws, r, 3, cd["vai_tro"])
        _cell(ws, r, 4, cd["ty_trong"], fmt=PCT, kind="input", align="center")
        _cell(ws, r, 5, "=" + budget_ref + "*D" + str(r), fmt=VND, kind="formula")
        _cell(ws, r, 6, "=E" + str(r) + "/30", fmt=VND, kind="formula")
        _cell(ws, r, 7, "=COUNTIF('2. Bộ từ khoá'!$B:$B,A" + str(r) + ")",
              kind="formula", align="center")
        cd_rows.append(r)
        r += 1
    cd_last = r - 1

    _cell(ws, r, 1, "Tổng", bold=True, fill=TOTAL)
    _cell(ws, r, 2, None, fill=TOTAL)
    _cell(ws, r, 3, None, fill=TOTAL)
    _cell(ws, r, 4, "=SUM(D%d:D%d)" % (cd_first, cd_last), fmt=PCT, bold=True, fill=TOTAL, align="center")
    _cell(ws, r, 5, "=SUM(E%d:E%d)" % (cd_first, cd_last), fmt=VND, bold=True, fill=TOTAL)
    _cell(ws, r, 6, "=SUM(F%d:F%d)" % (cd_first, cd_last), fmt=VND, bold=True, fill=TOTAL)
    _cell(ws, r, 7, "=SUM(G%d:G%d)" % (cd_first, cd_last), bold=True, fill=TOTAL, align="center")
    r += 2

    r = _header(ws, r, ["Cấp ưu tiên", "Tỷ trọng ngân sách", "Ngân sách/tháng (VND)", "Quy tắc đấu thầu"])
    for pr in z["cap_uu_tien"]:
        _cell(ws, r, 1, pr["cap"], bold=True, align="center")
        _cell(ws, r, 2, "=SUMIF($B$%d:$B$%d,A%d,$D$%d:$D$%d)" % (cd_first, cd_last, r, cd_first, cd_last),
              fmt=PCT, kind="formula", align="center")
        _cell(ws, r, 3, "=SUMIF($B$%d:$B$%d,A%d,$E$%d:$E$%d)" % (cd_first, cd_last, r, cd_first, cd_last),
              fmt=VND, kind="formula")
        _cell(ws, r, 4, pr["quy_tac"])
        r += 1
    r += 1

    # ---- C. mô hình phễu
    r = _title(ws, r, "C. MÔ HÌNH PHỄU: TỪ NGÂN SÁCH ĐẾN DOANH THU", 12)
    r = _header(ws, r, ["Chỉ số", "Giá trị", "Ghi chú / nguồn số"])
    p = z["phieu"]
    inputs = [("cpc", "CPC trung bình (VND)", VND),
              ("click_lead", "Tỷ lệ click -> lead", PCT1),
              ("lead_lich", "Tỷ lệ lead -> đặt lịch", PCT1),
              ("show_up", "Tỷ lệ đến khám (show-up)", PCT1),
              ("chot", "Tỷ lệ chốt ca", PCT1),
              ("gia_tri_ca", "Giá trị ca trung bình (VND)", VND)]
    ref = {}
    for key, label, fmt in inputs:
        _cell(ws, r, 1, label, bold=True)
        _cell(ws, r, 2, p[key], fmt=fmt, kind="input", align="center")
        _cell(ws, r, 3, p.get("ghi_chu", {}).get(key, ""))
        ref[key] = "B" + str(r)
        r += 1
    r += 1

    def outrow(label, formula, fmt, note):
        nonlocal r
        _cell(ws, r, 1, label, bold=True)
        _cell(ws, r, 2, formula, fmt=fmt, kind="formula", align="center")
        _cell(ws, r, 3, note)
        cur = "B" + str(r)
        r += 1
        return cur

    click = outrow("Click", "=IF(%s=0,0,%s/%s)" % (ref["cpc"], budget_ref, ref["cpc"]), NUM0, "Công thức")
    lead = outrow("Lead", "=%s*%s" % (click, ref["click_lead"]), NUM0,
                  "Lead = form + cuộc gọi >= 60 giây + chat Zalo có SĐT")
    lich = outrow("Lịch hẹn", "=%s*%s" % (lead, ref["lead_lich"]), NUM0, "Công thức")
    kham = outrow("Lượt đến khám", "=%s*%s" % (lich, ref["show_up"]), NUM0, "Công thức")
    chot = outrow("Ca chốt", "=%s*%s" % (kham, ref["chot"]), NUM0, "Công thức")
    doanh_thu = outrow("Doanh thu (VND)", "=%s*%s" % (chot, ref["gia_tri_ca"]), VND, "Công thức")
    cpl = outrow("CPL (VND)", "=IF(%s=0,0,%s/%s)" % (lead, budget_ref, lead), VND, "Công thức")
    cp_case = outrow("Chi phí / ca chốt (VND)", "=IF(%s=0,0,%s/%s)" % (chot, budget_ref, chot), VND,
                     "Chỉ số quyết định ngân sách, không phải CPL")
    outrow("ROAS", "=IF(%s=0,0,%s/%s)" % (budget_ref, doanh_thu, budget_ref), NUM2, "Công thức")
    r += 1

    # ---- D. nguyên tắc vận hành
    r = _title(ws, r, "D. NGUYÊN TẮC VẬN HÀNH", 12)
    for i, rule in enumerate(z["nguyen_tac"], start=1):
        _cell(ws, r, 1, "%d. %s" % (i, rule), wrap=False)
        r += 1

    return {"cd_rows": cd_rows, "cd_first": cd_first, "cd_last": cd_last,
            "cpl": cpl, "cp_case": cp_case, "gia_tri_ca": ref["gia_tri_ca"],
            "show_up": ref["show_up"], "chot": ref["chot"], "lead_lich": ref["lead_lich"]}


# ========================================================== 2. Bộ từ khoá
KW_COLS = [
    ("stt", "STT", 6),
    ("chien_dich", "Chiến dịch", 30),
    ("nhom_qc", "Nhóm quảng cáo", 24),
    ("tu_khoa", "Từ khoá", 38),
    ("khop", "Kiểu khớp", 11),
    ("uu_tien", "Ưu tiên", 9),
    ("chang", "Chặng hành trình", 13),
    ("chan_dung", "Chân dung KH", 12),
    ("rao_can", "Rào cản chốt", 13),
    ("canh_tranh", "Cạnh tranh dự kiến", 13),
    ("win", "Cổng WIN (/3)", 10),
    ("lp", "Landing page", 12),
    ("thong_diep", "Thông điệp / CTA chính", 40),
    ("volume", "Lượng tìm kiếm/tháng (Keyword Planner)", 15),
    ("cpc", "CPC đề xuất VND (Keyword Planner)", 15),
    ("ghi_chu", "Ghi chú vận hành", 44),
]
CENTER_KEYS = {"stt", "khop", "uu_tien", "chang", "chan_dung", "canh_tranh",
               "win", "lp", "volume", "cpc"}


def sheet_tu_khoa(wb, z):
    ws = wb.create_sheet("2. Bộ từ khoá")
    _widths(ws, dict((get_column_letter(i), w) for i, (_, _, w) in enumerate(KW_COLS, start=1)))
    _header(ws, 1, [lb for _, lb, _ in KW_COLS])

    r = 2
    for i, kw in enumerate(z["tu_khoa"], start=1):
        bg = WHITE if i % 2 else ZEBRA
        row = dict(kw)
        row["stt"] = i
        for ci, (key, _, _) in enumerate(KW_COLS, start=1):
            v = row.get(key, "")
            _cell(ws, r, ci, v if v != "" else None, fill=bg,
                  align="center" if key in CENTER_KEYS else None,
                  fmt=VND if key == "cpc" else None,
                  bold=(key == "tu_khoa"))
        r += 1

    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(KW_COLS)), r - 1)
    ws.freeze_panes = "E2"
    return r - 2


# ==================================================== 3. Từ khoá phủ định
def sheet_phu_dinh(wb, z):
    ws = wb.create_sheet("3. Từ khoá phủ định")
    _widths(ws, {"A": 34, "B": 30, "C": 12, "D": 46, "E": 50})
    r = _header(ws, 1, ["Nhóm", "Từ khoá phủ định", "Kiểu khớp", "Áp dụng ở", "Lý do"])

    r = _band(ws, r, "A. DANH SÁCH CHUNG - cấp tài khoản: " + z["zone"]["ds_phu_dinh_chung"], 5)
    for n in z["phu_dinh_chung"]:
        _cell(ws, r, 1, n["nhom"])
        _cell(ws, r, 2, n["tu"], bold=True)
        _cell(ws, r, 3, n["khop"], align="center")
        _cell(ws, r, 4, n["ap_dung"])
        _cell(ws, r, 5, n["ly_do"])
        r += 1
    r += 1

    r = _band(ws, r, "B. PHỦ ĐỊNH CHÉO - điều hướng mỗi truy vấn về đúng 1 chiến dịch", 5)
    for n in z["phu_dinh_cheo"]:
        _cell(ws, r, 1, n["nhom"])
        _cell(ws, r, 2, n["tu"], bold=True)
        _cell(ws, r, 3, n["khop"], align="center")
        _cell(ws, r, 4, n["ap_dung"])
        _cell(ws, r, 5, n["ly_do"])
        r += 1
    r += 1

    r = _band(ws, r, "C. QUY TRÌNH CẬP NHẬT", 5)
    for line in z["quy_trinh_phu_dinh"]:
        _cell(ws, r, 1, line, wrap=False)
        r += 1

    ws.freeze_panes = "A2"


# ====================================================== 4. Hành trình KH
JOURNEY_COLS = [
    ("chang", "Chặng (S1-S6)", 26),
    ("chan_dung", "Chân dung KH chính", 20),
    ("tam_ly", "Tâm lý & insight", 38),
    ("rao_can", "Rào cản chốt", 18),
    ("chien_dich", "Chiến dịch", 30),
    ("tu_khoa", "Cụm truy vấn đại diện", 30),
    ("thong_diep", "Thông điệp quảng cáo", 38),
    ("bang_chung", "Bằng chứng bắt buộc", 34),
    ("lp", "Landing page", 14),
    ("chuyen_doi", "Chuyển đổi đo lường", 26),
    ("sau_lead", "Hành động sau lead (Caresoft)", 34),
    ("kpi", "KPI chặng", 22),
    ("tiep_theo", "Bước tiếp theo", 26),
]


def sheet_hanh_trinh(wb, z):
    ws = wb.create_sheet("4. Hành trình KH")
    _widths(ws, dict((get_column_letter(i), w) for i, (_, _, w) in enumerate(JOURNEY_COLS, start=1)))
    r = _header(ws, 1, [lb for _, lb, _ in JOURNEY_COLS])

    for i, s in enumerate(z["hanh_trinh"]):
        bg = WHITE if i % 2 == 0 else ZEBRA
        for ci, (key, _, _) in enumerate(JOURNEY_COLS, start=1):
            _cell(ws, r, ci, s.get(key) or None, fill=bg,
                  bold=(key == "chang"),
                  align="center" if key in ("lp", "chan_dung") else None)
        r += 1
    r += 1

    r = _band(ws, r, "B. LANDING PAGE - 1 nhóm = 1 dịch vụ x 1 chặng x 1 intent = 1 trang", 8)
    r = _header(ws, r, ["Landing page", "Loại khung", "Chặng", "Chân dung KH",
                        "Rào cản gỡ chính", "Nội dung bắt buộc", "CTA chính",
                        "Micro-conversion trước form"])
    for lp in z["landing"]:
        _cell(ws, r, 1, lp["ma_ten"], bold=True)
        _cell(ws, r, 2, lp.get("loai", ""), align="center")
        _cell(ws, r, 3, lp.get("chang", ""), align="center")
        _cell(ws, r, 4, lp.get("chan_dung", ""), align="center")
        _cell(ws, r, 5, lp.get("rao_can", ""))
        _cell(ws, r, 6, lp["noi_dung"])
        _cell(ws, r, 7, lp["cta"])
        _cell(ws, r, 8, lp.get("micro", ""))
        r += 1
    r += 1

    if z.get("hook_format"):
        r = _band(ws, r, "C. NGÂN HÀNG HOOK x FORMAT (FB / TikTok) - cùng chân dung, cùng chặng với Search", 5)
        r = _header(ws, r, ["Chặng", "Chân dung KH", "Hook ưu tiên", "Format khả thi - hiệu quả",
                            "Mục tiêu thống nhất với Search"])
        for h in z["hook_format"]:
            _cell(ws, r, 1, h["chang"], bold=True, align="center")
            _cell(ws, r, 2, h["chan_dung"], align="center")
            _cell(ws, r, 3, h["hook"])
            _cell(ws, r, 4, h["format"])
            _cell(ws, r, 5, h["muc_tieu"])
            r += 1

    ws.freeze_panes = "B2"


# ================================================= 5. Giá thầu & Tối ưu
def sheet_gia_thau(wb, z, a):
    ws = wb.create_sheet("5. Giá thầu & Tối ưu")
    _widths(ws, {"A": 36, "B": 17, "C": 18, "D": 18, "E": 14, "F": 11, "G": 36, "H": 36})
    r = 1
    r = _title(ws, r, "A. CHẤM ĐIỂM ƯU TIÊN CHIẾN DỊCH", 12)
    r = _note(ws, r, "Điểm = Ý định mua x Giá trị ca x Khả năng chốt (thang 1-5). "
                     "Điểm ban đầu là đánh giá chuyên môn; cập nhật hằng tháng bằng tỷ lệ chốt thực tế từ Caresoft.")
    r = _header(ws, r, ["Chiến dịch", "Ý định mua (1-5)", "Giá trị ca (1-5)", "Khả năng chốt (1-5)",
                        "Điểm ưu tiên", "Xếp hạng", "Chiến lược giá thầu khởi đầu", "Chỉ tiêu cạnh tranh"])
    first = r
    last = first + len(z["chien_dich"]) - 1
    for cd, src in zip(z["chien_dich"], a["cd_rows"]):
        d = cd["diem"]
        _cell(ws, r, 1, "='1. Tổng quan'!A" + str(src), kind="formula", bold=True)
        _cell(ws, r, 2, d["y_dinh"], kind="input", align="center")
        _cell(ws, r, 3, d["gia_tri"], kind="input", align="center")
        _cell(ws, r, 4, d["kha_nang_chot"], kind="input", align="center")
        _cell(ws, r, 5, "=B%d*C%d*D%d" % (r, r, r), kind="formula", align="center")
        _cell(ws, r, 6, "=RANK(E%d,$E$%d:$E$%d)" % (r, first, last), kind="formula", align="center")
        _cell(ws, r, 7, cd["bid"])
        _cell(ws, r, 8, cd["chi_tieu"])
        r += 1
    r += 1

    r = _title(ws, r, "B. NGƯỠNG QUẢN TRỊ GIÁ THẦU", 12)
    r = _header(ws, r, ["Chỉ số", "Giá trị (VND)", "Ghi chú"])
    _cell(ws, r, 1, "CPL mục tiêu", bold=True)
    _cell(ws, r, 2, "='1. Tổng quan'!" + a["cpl"], fmt=VND, kind="formula", align="center")
    _cell(ws, r, 3, "Lấy từ mô hình phễu tab 1")
    cpl_row = r
    r += 1
    _cell(ws, r, 1, "Chi phí / ca chốt mục tiêu", bold=True)
    _cell(ws, r, 2, "='1. Tổng quan'!" + a["cp_case"], fmt=VND, kind="formula", align="center")
    _cell(ws, r, 3, "Lấy từ mô hình phễu tab 1")
    case_row = r
    r += 1
    _cell(ws, r, 1, "Giá trị gán cho 1 lịch hẹn", bold=True)
    _cell(ws, r, 2, "='1. Tổng quan'!%s*'1. Tổng quan'!%s*'1. Tổng quan'!%s"
          % (a["gia_tri_ca"], a["show_up"], a["chot"]), fmt=VND, kind="formula", align="center")
    _cell(ws, r, 3, "Dùng khi import chuyển đổi offline (Pha 3)")
    lich_row = r
    r += 1
    _cell(ws, r, 1, "Giá trị gán cho 1 lead", bold=True)
    _cell(ws, r, 2, "=B%d*'1. Tổng quan'!%s" % (lich_row, a["lead_lich"]),
          fmt=VND, kind="formula", align="center")
    _cell(ws, r, 3, "Dùng khi import chuyển đổi offline (Pha 3)")
    r += 2

    r = _header(ws, r, ["Quy tắc", "Hệ số", "Ngưỡng (VND)", "Hành động"])
    for rule in z["nguong"]:
        _cell(ws, r, 1, rule["quy_tac"], bold=True)
        _cell(ws, r, 2, rule["he_so"], kind="input", align="center")
        base = "$B$" + str(case_row if rule["moc"] == "ca_chot" else cpl_row)
        _cell(ws, r, 3, "=B%d*%s" % (r, base), fmt=VND, kind="formula", align="center")
        _cell(ws, r, 4, rule["hanh_dong"])
        r += 1
    r += 1

    r = _title(ws, r, "C. LỘ TRÌNH CHIẾN LƯỢC GIÁ THẦU 3 PHA", 12)
    r = _header(ws, r, ["Pha", "Thời gian", "Điều kiện chuyển pha", "Chiến lược giá thầu",
                        "Chuyển đổi tối ưu", "Việc bắt buộc"])
    for p in z["lo_trinh"]:
        _cell(ws, r, 1, p["pha"], bold=True)
        _cell(ws, r, 2, p["thoi_gian"], align="center")
        _cell(ws, r, 3, p["dieu_kien"])
        _cell(ws, r, 4, p["chien_luoc"])
        _cell(ws, r, 5, p["chuyen_doi"])
        _cell(ws, r, 6, p["bat_buoc"])
        r += 1
    r += 1

    r = _title(ws, r, "D. QUY TẮC CẠNH TRANH & VẬN HÀNH", 12)
    r = _header(ws, r, ["Hạng mục", "Quy tắc"])
    for q in z["quy_tac_canh_tranh"]:
        _cell(ws, r, 1, q["hang_muc"], bold=True)
        _cell(ws, r, 2, q["quy_tac"])
        r += 1
    r += 1

    r = _title(ws, r, "E. KẾ HOẠCH A/B TEST", 12)
    r = _header(ws, r, ["Test", "Biến thể A", "Biến thể B", "Chỉ số quyết định", "Thời gian tối thiểu"])
    for t in z["ab_test"]:
        _cell(ws, r, 1, t["test"], bold=True)
        _cell(ws, r, 2, t["a"])
        _cell(ws, r, 3, t["b"])
        _cell(ws, r, 4, t["chi_so"])
        _cell(ws, r, 5, t["thoi_gian"], align="center")
        r += 1
    r += 1

    r = _title(ws, r, "F. RỦI RO VÀ CÁCH KIỂM SOÁT", 12)
    r = _header(ws, r, ["Rủi ro", "Cách kiểm soát"])
    for k in z["rui_ro"]:
        _cell(ws, r, 1, k["rui_ro"], bold=True)
        _cell(ws, r, 2, k["kiem_soat"])
        r += 1


# ================================================================ kiểm tra
REQUIRED = ["zone", "chan_dung", "chien_dich", "cap_uu_tien", "phieu", "nguyen_tac",
            "tu_khoa", "phu_dinh_chung", "phu_dinh_cheo", "quy_trinh_phu_dinh",
            "hanh_trinh", "landing", "nguong", "lo_trinh", "quy_tac_canh_tranh",
            "ab_test", "rui_ro"]

CAM = ["tốt nhất", "số 1", "số một", "không đau 100", "cam kết thành công",
       "khỏi 100", "đẹp tuyệt đối", "vĩnh viễn", "duy nhất"]


def validate(z, path):
    errs = []
    for k in REQUIRED:
        if k not in z:
            errs.append("thiếu khoá bắt buộc: " + k)
    if errs:
        raise SystemExit("[LỖI] " + path + "\n  - " + "\n  - ".join(errs))

    tong = round(sum(c["ty_trong"] for c in z["chien_dich"]), 6)
    if abs(tong - 1.0) > 1e-6:
        errs.append("tổng tỷ trọng ngân sách = %.4f, phải bằng 1.0" % tong)

    ten_cd = set(c["ten"] for c in z["chien_dich"])
    ma_lp = set(lp["ma_ten"].split(":")[0].strip() for lp in z["landing"])
    ma_cd = set(p["ma"] for p in z["chan_dung"])
    cap = set(p["cap"] for p in z["cap_uu_tien"])

    for c in z["chien_dich"]:
        if c["uu_tien"] not in cap:
            errs.append("chiến dịch '%s': cấp ưu tiên '%s' chưa khai báo ở cap_uu_tien"
                        % (c["ten"], c["uu_tien"]))

    seen = set()
    for i, kw in enumerate(z["tu_khoa"], start=1):
        tag = "từ khoá #%d '%s'" % (i, kw.get("tu_khoa", "?"))
        if kw.get("chien_dich") not in ten_cd:
            errs.append("%s: chiến dịch '%s' không có trong danh sách chiến dịch"
                        % (tag, kw.get("chien_dich")))
        if kw.get("lp") and kw["lp"] not in ma_lp:
            errs.append("%s: landing '%s' chưa khai báo ở mục landing" % (tag, kw["lp"]))
        if kw.get("chan_dung") and kw["chan_dung"] not in ma_cd:
            errs.append("%s: chân dung '%s' chưa khai báo ở mục chan_dung" % (tag, kw["chan_dung"]))
        if kw.get("khop") == "Broad":
            errs.append("%s: Broad match bị cấm ở Pha 1 - xem nguyên tắc vận hành" % tag)
        low = (kw.get("thong_diep", "") or "").lower()
        for w in CAM:
            if w in low:
                errs.append("%s: thông điệp chứa từ cấm quảng cáo y tế '%s'" % (tag, w))
        key = ((kw.get("tu_khoa") or "").strip().lower(), kw.get("khop"))
        if key in seen:
            errs.append("trùng từ khoá x kiểu khớp: %s (%s)" % key)
        seen.add(key)

    if errs:
        raise SystemExit("[LỖI] " + path + "\n  - " + "\n  - ".join(errs))


def build(path, outdir):
    with open(path, encoding="utf-8") as f:
        z = json.load(f)
    validate(z, path)

    wb = Workbook()
    wb.remove(wb.active)
    anchors = sheet_tong_quan(wb, z)
    n_kw = sheet_tu_khoa(wb, z)
    sheet_phu_dinh(wb, z)
    sheet_hanh_trinh(wb, z)
    sheet_gia_thau(wb, z, anchors)
    wb.active = 0

    name = z["zone"].get("ten_file") or ("NKP _ Google Ads _ Bộ từ khoá " + z["zone"]["ten"])
    os.makedirs(outdir, exist_ok=True)
    dest = os.path.join(outdir, name + ".xlsx")
    wb.save(dest)
    return dest, n_kw, len(z["chien_dich"]), len(z["chan_dung"])


def main():
    ap = argparse.ArgumentParser(
        description="Sinh bộ từ khoá Google Ads theo ZONE cho Nha khoa Paris")
    ap.add_argument("configs", nargs="+", help="file JSON cấu hình ZONE (hỗ trợ wildcard)")
    ap.add_argument("-o", "--outdir", default="out", help="thư mục xuất (mặc định: out/)")
    args = ap.parse_args()

    paths = []
    for c in args.configs:
        paths.extend(sorted(glob.glob(c)) or [c])

    for p in paths:
        dest, n_kw, n_cd, n_persona = build(p, args.outdir)
        print("[OK] %s\n     -> %s\n     %d chân dung KH | %d chiến dịch | %d từ khoá"
              % (p, dest, n_persona, n_cd, n_kw))


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except AttributeError:
        pass
    main()
